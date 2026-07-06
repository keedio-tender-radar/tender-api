import io
import types

import openpyxl

from tender_api.services import xlsxmodels

_TENDER = types.SimpleNamespace(
    title="Servicios de soporte X",
    budget_amount=100000.0,
    currency="EUR",
    source_id="EXP-2026-1",
    source="placsp",
    buyer="Ayuntamiento de Ejemplo",
    deadline="2026-09-01",
    cpv=["72000000"],
    summary="Objeto del contrato de ejemplo.",
)
_SCORE = types.SimpleNamespace(
    total=80,
    recommendation="go",
    factors=[{"kind": "negative", "message": "Plazo ajustado"}],
    hard_rules=[],
    breakdown={},
)
_DRAFTS = [
    types.SimpleNamespace(
        kind="matriz_cumplimiento",
        content="| Requisito | Cumple | Evidencia |\n|---|---|---|\n| API REST | Sí | ok |",
    ),
    types.SimpleNamespace(
        kind="documentos_requeridos",
        content="- Declaración responsable\n- Solvencia técnica: 3 proyectos\n- Aval del 5%",
    ),
]


def test_build_plan_agil():
    data = xlsxmodels.build_plan_agil(_TENDER, _SCORE, _DRAFTS, ["Jefe", "Dev", "QA"], 6, 45.0, 0.2)
    assert data[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == [
        "HOJA1_Requisitos",
        "HOJA2_Supuestos",
        "HOJA3_Costes_Base",
        "HOJA4_Escenarios",
        "HOJA5_Resumen",
    ]
    # el motor de supuestos trae la duración
    assert wb["HOJA2_Supuestos"].cell(row=3, column=2).value == 6


def test_build_plan_exhaustivo():
    data = xlsxmodels.build_plan_exhaustivo(
        _TENDER, _SCORE, _DRAFTS, ["Jefe", "Dev", "QA"], 6, 45.0, 0.2
    )
    assert data[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for sheet in (
        "Portada",
        "00_RESUMEN_GO_NO_GO",
        "01_DATOS_EXPEDIENTE",
        "02_JORNADAS_PLIEGO",
        "04_COSTE_POR_PERFIL",
        "05_COSTES_INDIRECTOS",
        "06_PRECIO_Y_MARGEN",
        "07_SIM_PUNTOS_PRECIO",
        "08_PUNT_TECNICA",
        "09_CHECKLIST_DOCS",
        "10_RIESGOS_CRITICOS",
        "BASE_Tarifas",
    ):
        assert sheet in wb.sheetnames, sheet
    # el expediente se rellena con el organismo real
    assert wb["01_DATOS_EXPEDIENTE"].cell(row=3, column=2).value == "Ayuntamiento de Ejemplo"
    # el riesgo del score aparece
    assert wb["10_RIESGOS_CRITICOS"].cell(row=4, column=1).value == "Plazo ajustado"
