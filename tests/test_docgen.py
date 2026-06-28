from tests.conftest import full_breakdown, make_tender


def _setup(client):
    t = make_tender(client)
    bd = full_breakdown()
    client.put(
        f"/api/tenders/{t['id']}/score",
        json={"total": sum(bd.values()), "breakdown": bd, "recommendation": "go",
              "summary": "Encaje alto."},
    )
    # inyecta borradores directamente vía generate? usamos notes-less: crea docs por la BD
    return t


def test_docx_and_pdf_download(client):
    from tender_api.database import SessionLocal
    from tender_api.models import GeneratedDocument
    t = _setup(client)
    # crea un par de borradores con tabla markdown
    with SessionLocal() as s:
        s.add(GeneratedDocument(tender_id=t["id"], kind="go_no_go", title="Go",
                                content="# Go\n\nTexto."))
        s.add(GeneratedDocument(
            tender_id=t["id"], kind="matriz", title="Matriz",
            content="# Matriz\n\n| Requisito | Cumple |\n|---|---|\n| API REST | Sí |",
        ))
        s.commit()
    docx = client.get(f"/api/tenders/{t['id']}/package.docx")
    assert docx.status_code == 200
    assert "wordprocessingml" in docx.headers["content-type"]
    assert docx.content[:2] == b"PK"  # zip/docx
    pdf = client.get(f"/api/tenders/{t['id']}/package.pdf")
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content[:4] == b"%PDF"


def test_plan_xlsx_download(client):
    t = make_tender(client)
    r = client.get(f"/api/tenders/{t['id']}/plan.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"  # xlsx es zip
    # comprueba las 3 hojas
    import io

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Requerimientos", "Cronograma", "Resumen de costes"]
