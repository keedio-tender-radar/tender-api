"""Módulo de planificación/estimación por licitación, exportado a Excel (.xlsx).

Replica la estructura de un cronograma de oferta tipo Keedio: Requerimientos (horas por perfil),
Cronograma por semanas (Gantt) y Resumen de costes (horas × tarifa + margen → precio). Plantilla
lista para que el equipo la rellene; alimentada por el perfil (equipo, duración, tarifa, margen).
"""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from io import BytesIO

BRAND_HEX = "5B94FF"
_DEFAULT_TEAM = ["Sistemas", "UX", "Front", "Back", "Data"]


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="D9DEE8")
    return {
        "hdr_font": Font(bold=True, color="FFFFFF", size=10),
        "hdr_fill": PatternFill("solid", fgColor=BRAND_HEX),
        "bold": Font(bold=True),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "wrap": Alignment(vertical="top", wrap_text=True),
    }


def build_project_plan(tender, score, team, months, rate, margin, drafts=None) -> bytes:
    from openpyxl import Workbook

    from tender_api.services import estimate

    team = [t for t in (team or []) if t] or _DEFAULT_TEAM
    months = int(months or 6)
    rate = float(rate or 45.0)
    margin = float(margin if margin is not None else 0.2)
    reqs = estimate.requirements_from_drafts(drafts)
    st = _styles()

    wb = Workbook()
    _sheet_portada(wb, wb.active, tender, score, st)
    _sheet_requerimientos(wb, wb.create_sheet("Requerimientos"), tender, team, st, reqs)
    _sheet_cronograma(wb, tender, months, st)
    _sheet_costes(wb, rate, margin, st)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_portada(wb, ws, tender, score, st) -> None:
    """Hoja de portada del plan: marca, título, datos clave, valoración y sello."""
    from openpyxl.styles import Alignment, Font

    ws.title = "Portada"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 92

    ws["B2"] = "KEEDIO"
    ws["B2"].font = Font(bold=True, size=28, color=BRAND_HEX)
    ws["B4"] = "PAQUETE DE OFERTA · PLAN DE PROYECTO"
    ws["B4"].font = Font(bold=True, size=11, color=BRAND_HEX)

    ws["B6"] = tender.title or "(sin título)"
    ws["B6"].font = Font(bold=True, size=20)
    ws["B6"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[6].height = 48

    ws["B8"] = f"Expediente {tender.source_id}  ·  Fuente {tender.source}"
    ws["B9"] = f"Presupuesto de licitación: {tender.budget_amount or 's/d'} {tender.currency}"
    if score:
        ws["B11"] = f"Valoración Go/No-Go: {score.total}/100 · {score.recommendation.upper()}"
        ws["B11"].font = Font(bold=True, size=13, color=BRAND_HEX)
    ws["B13"] = "Contenido: Requerimientos · Cronograma (Gantt) · Resumen de costes"
    ws["B13"].font = Font(italic=True, size=10, color="6E7686")
    ws["B15"] = (
        f"Generado el {datetime.now(UTC).date().isoformat()} · Keedio Tender Radar · Confidencial"
    )
    ws["B15"].font = Font(size=9, color="999999")


def _sheet_requerimientos(wb, ws, tender, team, st, reqs=None) -> None:
    from openpyxl.utils import get_column_letter

    from tender_api.services import estimate

    reqs = reqs or []
    ws.title = "Requerimientos"
    ws["A1"] = f"Keedio · Estimación — {tender.title or ''} ({tender.source_id})"
    ws["A1"].font = st["bold"]

    base = ["ID-Orden", "PT", "Requisito", "Descripción funcional", "Dependencias",
            "Perfiles involucrados"]
    hour_cols = [f"Horas {r}" for r in team]
    headers = base + hour_cols + ["Horas total"]
    hr = 3
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(hr, c, h)
        cell.font = st["hdr_font"]
        cell.fill = st["hdr_fill"]
        cell.alignment = st["center"]
        cell.border = st["border"]

    first_hour = len(base) + 1
    last_hour = first_hour + len(team) - 1
    total_col = last_hour + 1
    rows = max(len(reqs), 8)  # filas prerellenadas o, al menos, plantilla vacía
    for i in range(rows):
        r = hr + 1 + i
        ws.cell(r, 1, f"REQ-{i + 1:02d}")
        if i < len(reqs):
            ws.cell(r, 3, reqs[i][0]).alignment = st["wrap"]  # Requisito
            if reqs[i][1]:
                ws.cell(r, 4, reqs[i][1]).alignment = st["wrap"]  # Descripción/Evidencia
            # Horas sugeridas por rol según complejidad (editable).
            hours = estimate.suggest_hours(reqs[i][0], team)
            for j, role in enumerate(team):
                if hours.get(role):
                    ws.cell(r, first_hour + j, hours[role])
        for c in range(1, total_col + 1):
            ws.cell(r, c).border = st["border"]
        fl, ll = get_column_letter(first_hour), get_column_letter(last_hour)
        ws.cell(r, total_col, f"=SUM({fl}{r}:{ll}{r})")
    # Fila de totales.
    tr = hr + 1 + rows
    ws.cell(tr, 1, "TOTAL").font = st["bold"]
    for c in range(first_hour, total_col + 1):
        col = get_column_letter(c)
        ws.cell(tr, c, f"=SUM({col}{hr + 1}:{col}{tr - 1})").font = st["bold"]
        ws.cell(tr, c).border = st["border"]
    ws._ktr_total_cell = f"'Requerimientos'!{get_column_letter(total_col)}{tr}"

    widths = [12, 6, 26, 34, 14, 20] + [9] * len(team) + [11]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A4"


def _sheet_cronograma(wb, tender, months, st) -> None:
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("Cronograma")
    weeks = min(months * 4, 24)
    start = datetime.now(UTC).date()
    ws["A1"] = "Cronograma (Gantt) — semanas"
    ws["A1"].font = st["bold"]

    hr = 3
    base = ["Paquete", "Reqs.", "Descripción"]
    for c, h in enumerate(base, start=1):
        cell = ws.cell(hr, c, h)
        cell.font, cell.fill, cell.alignment, cell.border = (
            st["hdr_font"], st["hdr_fill"], st["center"], st["border"]
        )
    for w in range(weeks):
        c = len(base) + 1 + w
        d = start + timedelta(weeks=w)
        cell = ws.cell(hr, c, f"S{w + 1:02d}\n{d.strftime('%d/%m')}")
        cell.font, cell.fill, cell.alignment, cell.border = (
            st["hdr_font"], st["hdr_fill"], st["center"], st["border"]
        )
        ws.column_dimensions[get_column_letter(c)].width = 7

    for i in range(8):  # filas-paquete (plantilla)
        r = hr + 1 + i
        ws.cell(r, 1, f"PT{i + 1}" if i < 4 else "")
        for c in range(1, len(base) + 1 + weeks + 1):
            ws.cell(r, c).border = st["border"]

    leg = hr + 1 + 9
    ws.cell(leg, 1, "Leyenda:").font = st["bold"]
    ws.cell(leg, 2, "x=inicio · f=Front · b=Back · d=Data · s=Sys · t=Test")
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 34
    ws.freeze_panes = "D4"


def _sheet_costes(wb, rate, margin, st) -> None:
    ws = wb.create_sheet("Resumen de costes")
    total = getattr(wb["Requerimientos"], "_ktr_total_cell", "0")
    rows = [
        ("Horas totales (estimadas)", f"={total}"),
        ("Coste por hora (€)", rate),
        ("Costes generales (%)", 0.04),
        ("Margen comercial (%)", margin),
        ("Coste base (€)", "=B1*B2"),
        ("Con costes generales (€)", "=B5*(1+B3)"),
        ("Precio total (€)", "=B6*(1+B4)"),
        ("Precio por hora (€)", "=IF(B1=0,0,B7/B1)"),
    ]
    for i, (label, val) in enumerate(rows, start=1):
        ws.cell(i, 1, label).font = st["bold"] if i in (7,) else None
        ws.cell(i, 2, val)
    ws.cell(7, 1).font = st["bold"]
    ws.cell(7, 2).font = st["bold"]
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 16
